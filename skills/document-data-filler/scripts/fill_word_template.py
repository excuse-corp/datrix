"""Fill a DOCX template without modifying the original upload."""

import re
import sys
from typing import Any, Iterable

from docx import Document

from common import (
    completed_or_partial,
    emit_result,
    error_result,
    load_query_result,
    output_path,
    parse_assignments,
    payload_from_argv,
    require_input_file,
    text_value,
)


_TABLE_CELL_RE = re.compile(r"^table\[(\d+)\]\.cell\[(\d+)[,\]](\d+)\]$")
_PARAGRAPH_RE = re.compile(r"^paragraph\[(\d+)\]$")
_PLACEHOLDER_RE = re.compile(r"^(?:\{\{\s*[^{}]+?\s*\}\}|<<\s*[^<>]+?\s*>>)$")


def _iter_table_paragraphs(table: Any) -> Iterable[Any]:
    for row in table.rows:
        for cell in row.cells:
            yield from cell.paragraphs
            for nested in cell.tables:
                yield from _iter_table_paragraphs(nested)


def _iter_all_paragraphs(document: Any) -> Iterable[Any]:
    yield from document.paragraphs
    for table in document.tables:
        yield from _iter_table_paragraphs(table)
    for section in document.sections:
        for container in (section.header, section.footer):
            yield from container.paragraphs
            for table in container.tables:
                yield from _iter_table_paragraphs(table)


def _replace_in_paragraph(paragraph: Any, target: str, replacement: str) -> int:
    count = 0
    while True:
        runs = list(paragraph.runs)
        combined = "".join(run.text or "" for run in runs)
        start = combined.find(target)
        if start < 0:
            return count
        end = start + len(target)

        offsets: list[tuple[int, int]] = []
        cursor = 0
        for run in runs:
            run_text = run.text or ""
            offsets.append((cursor, cursor + len(run_text)))
            cursor += len(run_text)

        start_index = next(
            (index for index, (_, stop) in enumerate(offsets) if start < stop),
            None,
        )
        end_index = next(
            (index for index, (_, stop) in enumerate(offsets) if end <= stop),
            None,
        )
        if start_index is None or end_index is None:
            return count

        start_local = start - offsets[start_index][0]
        end_local = end - offsets[end_index][0]
        if start_index == end_index:
            run_text = runs[start_index].text or ""
            runs[start_index].text = (
                run_text[:start_local] + replacement + run_text[end_local:]
            )
        else:
            first_text = runs[start_index].text or ""
            last_text = runs[end_index].text or ""
            runs[start_index].text = first_text[:start_local] + replacement
            for index in range(start_index + 1, end_index):
                runs[index].text = ""
            runs[end_index].text = last_text[end_local:]
        count += 1


def _placeholder_variants(target: str) -> list[str]:
    target = target.strip()
    if target.startswith("{{") and target.endswith("}}"):
        field = target[2:-2].strip()
    elif target.startswith("<<") and target.endswith(">>"):
        field = target[2:-2].strip()
    else:
        field = target
    return [f"{{{{{field}}}}}", f"<<{field}>>"] if field else []


def _fill_document(
    document: Any, assignments: list[tuple[str, Any]]
) -> tuple[list[str], list[str], list[str]]:
    filled: list[str] = []
    unresolved: list[str] = []
    warnings: list[str] = []

    for target, value in assignments:
        target = target.strip()
        table_match = _TABLE_CELL_RE.fullmatch(target)
        if table_match:
            table_number, row_number, column_number = map(int, table_match.groups())
            if table_number > len(document.tables):
                unresolved.append(target)
                warnings.append(f"Table not found: {target}")
                continue
            table = document.tables[table_number - 1]
            if row_number > len(table.rows) or column_number > len(table.columns):
                unresolved.append(target)
                warnings.append(f"Table cell out of range: {target}")
                continue
            table.cell(row_number - 1, column_number - 1).text = text_value(value)
            filled.append(target)
            continue

        paragraph_match = _PARAGRAPH_RE.fullmatch(target)
        if paragraph_match:
            paragraph_number = int(paragraph_match.group(1))
            if paragraph_number < 1 or paragraph_number > len(document.paragraphs):
                unresolved.append(target)
                warnings.append(f"Paragraph not found: {target}")
                continue
            document.paragraphs[paragraph_number - 1].text = text_value(value)
            filled.append(target)
            continue

        if _PLACEHOLDER_RE.fullmatch(target) or not target.startswith("table["):
            replacement_count = 0
            for paragraph in _iter_all_paragraphs(document):
                for variant in _placeholder_variants(target):
                    replacement_count += _replace_in_paragraph(
                        paragraph, variant, text_value(value)
                    )
            if replacement_count:
                filled.append(target)
            else:
                unresolved.append(target)
                warnings.append(f"Placeholder not found: {target}")
            continue

        unresolved.append(target)
        warnings.append(f"Invalid Word target: {target}")

    return filled, unresolved, warnings


def _document_text(document: Any) -> str:
    parts = [paragraph.text for paragraph in _iter_all_paragraphs(document)]
    return "\n".join(parts)


def _validate_output(document: Any, assignments: list[tuple[str, Any]]) -> list[str]:
    errors: list[str] = []
    all_text = _document_text(document)
    for target, value in assignments:
        table_match = _TABLE_CELL_RE.fullmatch(target)
        if table_match:
            table_number, row_number, column_number = map(int, table_match.groups())
            if table_number <= len(document.tables):
                actual = document.tables[table_number - 1].cell(
                    row_number - 1, column_number - 1
                ).text
                if actual != text_value(value):
                    errors.append(f"Output value mismatch at {target}")
        elif not table_match and not _PARAGRAPH_RE.fullmatch(target) and any(
            variant in all_text for variant in _placeholder_variants(target)
        ):
            errors.append(f"Placeholder remains in output: {target}")
    return errors


def main(payload: dict[str, Any]) -> None:
    try:
        input_file = require_input_file(payload)
        if input_file.suffix.lower() != ".docx":
            emit_result(
                error_result(
                    "Only .docx templates are writable by this MVP.",
                    status="need_clarification",
                    code="unsupported_format",
                )
            )
            return

        query_result_text = load_query_result(payload)
        assignments, parse_errors = parse_assignments(payload)
        if parse_errors:
            result = error_result(
                "; ".join(parse_errors),
                status="need_clarification",
                code="invalid_assignments",
            )
            result["warnings"] = parse_errors
            emit_result(result)
            return
        if not assignments:
            emit_result(
                error_result(
                    "At least one explicit fill assignment is required.",
                    status="need_clarification",
                    code="missing_assignments",
                )
            )
            return

        output_file = output_path(input_file, payload)
        document = Document(str(output_file))
        filled, unresolved, warnings = _fill_document(document, assignments)
        document.save(output_file)

        reopened = Document(str(output_file))
        warnings.extend(_validate_output(reopened, assignments))
        result = completed_or_partial(
            output_file=output_file,
            filled=filled,
            unresolved=unresolved,
            warnings=warnings,
            query_result_text=query_result_text,
        )
        emit_result(result)
    except Exception as exc:
        emit_result(error_result(str(exc), code="fill_error"))


if __name__ == "__main__":
    try:
        main(payload_from_argv(sys.argv))
    except Exception as exc:
        emit_result(error_result(str(exc), code="invalid_arguments"))
