"""Fill an XLSX template without modifying the original upload."""

import re
import sys
from typing import Any

from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell

from common import (
    completed_or_partial,
    emit_result,
    error_result,
    load_query_result,
    output_path,
    parse_assignments,
    payload_from_argv,
    require_input_file,
    text_value,
)


_CELL_RE = re.compile(r"^[A-Za-z]{1,3}[1-9][0-9]*$")
_PLACEHOLDER_RE = re.compile(r"^(?:\{\{\s*[^{}]+?\s*\}\}|<<\s*[^<>]+?\s*>>)$")


def _split_cell_target(target: str, workbook: Any) -> tuple[Any, str] | None:
    if "!" in target:
        sheet_name, coordinate = target.rsplit("!", 1)
        sheet_name = sheet_name.strip().strip("'")
        coordinate = coordinate.strip().upper()
        if sheet_name in workbook.sheetnames and _CELL_RE.fullmatch(coordinate):
            return workbook[sheet_name], coordinate
        return None
    coordinate = target.strip().upper()
    if _CELL_RE.fullmatch(coordinate) and len(workbook.sheetnames) == 1:
        return workbook.active, coordinate
    return None


def _placeholder_variants(target: str) -> list[str]:
    target = target.strip()
    if target.startswith("{{") and target.endswith("}}"):
        field = target[2:-2].strip()
    elif target.startswith("<<") and target.endswith(">>"):
        field = target[2:-2].strip()
    else:
        field = target
    if not field:
        return []
    return [f"{{{{{field}}}}}", f"<<{field}>>"]


def _fill_workbook(
    workbook: Any, assignments: list[tuple[str, Any]]
) -> tuple[list[str], list[str], list[str]]:
    filled: list[str] = []
    unresolved: list[str] = []
    warnings: list[str] = []

    for target, value in assignments:
        target = target.strip()
        cell_target = _split_cell_target(target, workbook)
        if cell_target:
            worksheet, coordinate = cell_target
            cell = worksheet[coordinate]
            if isinstance(cell, MergedCell):
                unresolved.append(target)
                warnings.append(f"Cannot write non-anchor merged cell: {target}")
                continue
            if isinstance(cell.value, str) and cell.value.startswith("="):
                unresolved.append(target)
                warnings.append(f"Formula cell was not overwritten: {target}")
                continue
            cell.value = value
            filled.append(target)
            continue

        if _PLACEHOLDER_RE.fullmatch(target) or "!" not in target:
            variants = _placeholder_variants(target)
            replacement_count = 0
            for worksheet in workbook.worksheets:
                for row in worksheet.iter_rows():
                    for cell in row:
                        if isinstance(cell, MergedCell) or not isinstance(cell.value, str):
                            continue
                        if cell.value.startswith("="):
                            continue
                        original = cell.value
                        updated = original
                        for variant in variants:
                            updated = updated.replace(variant, text_value(value))
                        if updated != original:
                            cell.value = updated
                            replacement_count += 1
            if replacement_count:
                filled.append(target)
            else:
                unresolved.append(target)
                warnings.append(f"Placeholder not found: {target}")
            continue

        unresolved.append(target)
        warnings.append(f"Invalid worksheet/cell target: {target}")

    return filled, unresolved, warnings


def _validate_output(workbook: Any, assignments: list[tuple[str, Any]]) -> list[str]:
    errors: list[str] = []
    for target, value in assignments:
        cell_target = _split_cell_target(target, workbook)
        if cell_target:
            worksheet, coordinate = cell_target
            actual = worksheet[coordinate].value
            if actual != value and text_value(actual) != text_value(value):
                errors.append(
                    f"Output value mismatch at {target}: expected {value!r}, got {actual!r}"
                )
            continue
        variants = _placeholder_variants(target)
        if variants and any(
            isinstance(cell.value, str)
            and any(variant in cell.value for variant in variants)
            for worksheet in workbook.worksheets
            for row in worksheet.iter_rows()
            for cell in row
        ):
            errors.append(f"Placeholder remains in output: {target}")
    return errors


def main(payload: dict[str, Any]) -> None:
    try:
        input_file = require_input_file(payload)
        if input_file.suffix.lower() != ".xlsx":
            message = "Only .xlsx templates are writable by this MVP; convert .xls to .xlsx first."
            emit_result(
                error_result(
                    message,
                    status="need_clarification",
                    code="unsupported_format",
                )
            )
            return

        query_result_text = load_query_result(payload)
        assignments, parse_errors = parse_assignments(payload)
        if parse_errors:
            result = error_result(
                "; ".join(parse_errors),
                status="need_clarification",
                code="invalid_assignments",
            )
            result["warnings"] = parse_errors
            emit_result(result)
            return
        if not assignments:
            emit_result(
                error_result(
                    "At least one explicit fill assignment is required.",
                    status="need_clarification",
                    code="missing_assignments",
                )
            )
            return

        output_file = output_path(input_file, payload)
        workbook = load_workbook(output_file, data_only=False)
        filled, unresolved, warnings = _fill_workbook(workbook, assignments)
        workbook.save(output_file)

        reopened = load_workbook(output_file, data_only=False)
        validation_errors = _validate_output(reopened, assignments)
        if validation_errors:
            warnings.extend(validation_errors)
            unresolved.extend(
                target
                for target, _ in assignments
                if target not in unresolved and target not in filled
            )
        result = completed_or_partial(
            output_file=output_file,
            filled=filled,
            unresolved=unresolved,
            warnings=warnings,
            query_result_text=query_result_text,
        )
        emit_result(result)
    except Exception as exc:
        emit_result(error_result(str(exc), code="fill_error"))


if __name__ == "__main__":
    try:
        main(payload_from_argv(sys.argv))
    except Exception as exc:
        emit_result(error_result(str(exc), code="invalid_arguments"))
