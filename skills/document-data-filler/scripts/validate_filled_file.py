"""Validate explicit targets in an already generated DOCX or XLSX file."""

import sys

from common import (
    emit_result,
    error_result,
    parse_assignments,
    payload_from_argv,
    require_input_file,
)


def main(payload: dict) -> None:
    input_file = require_input_file(payload)
    assignments, parse_errors = parse_assignments(payload)
    if parse_errors:
        result = error_result(
            "; ".join(parse_errors),
            status="need_clarification",
            code="invalid_assignments",
        )
        result["warnings"] = parse_errors
        emit_result(result)
        return
    if not assignments:
        emit_result(
            error_result(
                "At least one explicit fill assignment is required.",
                status="need_clarification",
                code="missing_assignments",
            )
        )
        return

    if input_file.suffix.lower() == ".xlsx":
        from fill_excel_template import _validate_output
        from openpyxl import load_workbook

        document = load_workbook(input_file, data_only=False)
        errors = _validate_output(document, assignments)
    elif input_file.suffix.lower() == ".docx":
        from docx import Document
        from fill_word_template import _validate_output

        document = Document(str(input_file))
        errors = _validate_output(document, assignments)
    else:
        emit_result(
            error_result(
                "Only .xlsx and .docx files can be validated.",
                status="need_clarification",
                code="unsupported_format",
            )
        )
        return

    emit_result(
        {
            "status": "completed" if not errors else "partial",
            "file_path": str(input_file),
            "filename": input_file.name,
            "filled_fields": [target for target, _ in assignments]
            if not errors
            else [],
            "unresolved_fields": errors,
            "warnings": errors,
        }
    )


if __name__ == "__main__":
    try:
        main(payload_from_argv(sys.argv))
    except Exception as exc:
        emit_result(error_result(str(exc), code="validation_error"))
