from types import SimpleNamespace

import pytest

from dbgpt_serve.utils.auth import UserRequest


@pytest.mark.asyncio
async def test_attachment_parse_api_scopes_paths_and_returns_per_file_errors(
    tmp_path, monkeypatch
):
    from dbgpt_app.attachment import AttachmentParseService
    from dbgpt_app.openapi.api_v1 import attachment_api, python_upload_api

    monkeypatch.setattr(
        python_upload_api.CFG,
        "SYSTEM_APP",
        SimpleNamespace(work_dir=str(tmp_path)),
    )
    monkeypatch.setattr(
        attachment_api,
        "AttachmentParseService",
        lambda **kwargs: AttachmentParseService(
            storage_dir=tmp_path / "cache", **kwargs
        ),
    )
    upload_dir = tmp_path / "python_uploads" / "alice"
    upload_dir.mkdir(parents=True)
    good = upload_dir / "note.txt"
    good.write_text("hello attachment", encoding="utf-8")
    outside = tmp_path / "outside.txt"
    outside.write_text("do not read", encoding="utf-8")

    result = await attachment_api.parse_attachments(
        attachment_api.AttachmentParseRequest(
            file_paths=[str(good), str(outside)],
        ),
        UserRequest(user_id="alice"),
    )
    assert result.success is True
    assert [item["status"] for item in result.data] == ["completed", "failed"]
    assert "do not read" not in json_text(result.data[1])
    assert str(outside) not in json_text(result.data[1])


def json_text(value):
    import json

    return json.dumps(value, ensure_ascii=False)


def test_attachment_parse_route_is_mounted():
    from fastapi import FastAPI

    from dbgpt_app.dbgpt_server import mount_routers

    app = FastAPI()
    mount_routers(app)

    route_paths = {route.path for route in app.routes}
    assert "/api/v1/attachments/parse" in route_paths
    assert "/api/v1/attachments/parse/{job_id}" in route_paths


@pytest.mark.asyncio
async def test_attachment_parse_api_rejects_more_than_five_files():
    from dbgpt_app.openapi.api_v1 import attachment_api

    with pytest.raises(ValueError):
        attachment_api.AttachmentParseRequest(
            file_paths=[f"/tmp/{index}.txt" for index in range(6)]
        )


@pytest.mark.asyncio
async def test_attachment_parse_api_rejects_oversized_path_string():
    from dbgpt_app.openapi.api_v1 import attachment_api

    result = await attachment_api.parse_attachments(
        attachment_api.AttachmentParseRequest(file_paths=["x" * 4097]),
        UserRequest(user_id="alice"),
    )

    assert result.success is False
    assert result.err_code == "ATTACHMENT_PATH_INVALID"


@pytest.mark.asyncio
async def test_attachment_parse_api_rejects_misaligned_ids_before_path_lookup():
    from dbgpt_app.openapi.api_v1 import attachment_api

    result = await attachment_api.parse_attachments(
        attachment_api.AttachmentParseRequest(
            file_paths=["/tmp/one.txt"],
            attachment_ids=["att_one", "att_two"],
        ),
        UserRequest(user_id="alice"),
    )

    assert result.success is False
    assert result.err_code == "ATTACHMENT_ID_MISMATCH"


def test_attachment_parse_request_accepts_aligned_ids():
    from dbgpt_app.openapi.api_v1 import attachment_api

    request = attachment_api.AttachmentParseRequest(
        file_paths=["/tmp/note.txt"],
        attachment_ids=["att_test"],
    )
    assert request.attachment_ids == ["att_test"]
    assert request.async_parse is False


@pytest.mark.asyncio
async def test_attachment_parse_api_submits_async_job_and_scopes_status(
    tmp_path, monkeypatch
):
    import asyncio

    from dbgpt_app.attachment import AttachmentParseService
    from dbgpt_app.attachment import storage
    from dbgpt_app.attachment.jobs import AttachmentParseJobManager
    from dbgpt_app.attachment.storage import AttachmentChunkStore
    from dbgpt_app.openapi.api_v1 import attachment_api, python_upload_api

    monkeypatch.setattr(
        python_upload_api.CFG,
        "SYSTEM_APP",
        SimpleNamespace(work_dir=str(tmp_path)),
    )
    upload_dir = tmp_path / "python_uploads" / "alice"
    upload_dir.mkdir(parents=True)
    source = upload_dir / "async.txt"
    source.write_text("asynchronous attachment", encoding="utf-8")

    def parse_batch(*args, **kwargs):
        kwargs.pop("cache_namespace", None)
        return AttachmentParseService(storage_dir=tmp_path / "cache").parse_files(
            *args, **kwargs
        )

    manager = AttachmentParseJobManager(
        storage_dir=tmp_path / "jobs",
        max_workers=1,
        max_concurrent_jobs_per_user=1,
        parse_batch=parse_batch,
    )
    store = AttachmentChunkStore(storage_dir=tmp_path / "index")
    monkeypatch.setattr(attachment_api, "get_attachment_job_manager", lambda: manager)
    monkeypatch.setattr(storage, "get_attachment_chunk_store", lambda: store)

    try:
        submitted = await attachment_api.parse_attachments(
            attachment_api.AttachmentParseRequest(
                file_paths=[str(source)], async_parse=True
            ),
            UserRequest(user_id="alice"),
        )
        assert submitted.success is True
        assert submitted.data["job_id"].startswith("parse_job_")

        status_result = None
        for _ in range(80):
            status_result = await attachment_api.get_parse_job(
                submitted.data["job_id"], UserRequest(user_id="alice")
            )
            if status_result.data["status"] in {"completed", "failed"}:
                break
            await asyncio.sleep(0.025)
        assert status_result.data["status"] == "completed"
        assert status_result.data["items"][0]["text_chars"] > 0

        hidden_status = await attachment_api.get_parse_job(
            submitted.data["job_id"], UserRequest(user_id="bob")
        )
        assert hidden_status.success is False
        assert hidden_status.err_code == "ATTACHMENT_JOB_NOT_FOUND"
    finally:
        manager.shutdown()


def test_attachment_parse_job_manager_runs_and_persists(tmp_path, monkeypatch):
    import time

    from dbgpt_app.attachment.jobs import AttachmentParseJobManager
    from dbgpt_app.attachment.storage import AttachmentChunkStore
    from dbgpt_app.attachment import storage

    index = AttachmentChunkStore(storage_dir=tmp_path / "index")
    monkeypatch.setattr(storage, "get_attachment_chunk_store", lambda: index)

    upload_dir = tmp_path / "uploads"
    upload_dir.mkdir()
    source = upload_dir / "note.txt"
    source.write_text("background attachment", encoding="utf-8")
    manager = AttachmentParseJobManager(
        storage_dir=tmp_path / "jobs",
        max_workers=1,
        max_concurrent_jobs_per_user=1,
    )
    job = manager.submit(
        user_id="alice",
        file_paths=[source],
        attachment_ids=["att_background"],
        allowed_roots=[upload_dir],
    )
    for _ in range(40):
        current = manager.get(job.job_id, user_id="alice")
        if current.status in {"completed", "failed", "partial", "needs_ocr"}:
            break
        time.sleep(0.025)
    assert current.status == "completed"
    assert current.items[0]["attachment_id"] == "att_background"
    assert (tmp_path / "jobs" / "jobs.sqlite3").is_file()
    assert index.search(
        "alice", "background attachment", attachment_ids=["att_background"]
    )
    manager.shutdown()
    restarted = AttachmentParseJobManager(
        storage_dir=tmp_path / "jobs",
        max_workers=1,
        max_concurrent_jobs_per_user=1,
    )
    loaded = restarted.get(job.job_id, user_id="alice")
    assert loaded.status == "completed"
    assert loaded.items == current.items
    with pytest.raises(KeyError):
        restarted.get(job.job_id, user_id="bob")
    restarted.shutdown()


def test_attachment_job_concurrency_limit_is_shared_between_managers(
    tmp_path, monkeypatch
):
    import threading

    from dbgpt_app.attachment.jobs import (
        AttachmentJobLimitError,
        AttachmentParseJobManager,
    )
    from dbgpt_app.attachment import storage
    from dbgpt_app.attachment.storage import AttachmentChunkStore

    monkeypatch.setattr(
        storage,
        "get_attachment_chunk_store",
        lambda: AttachmentChunkStore(storage_dir=tmp_path / "index"),
    )

    upload_dir = tmp_path / "uploads"
    upload_dir.mkdir()
    first = upload_dir / "first.txt"
    second = upload_dir / "second.txt"
    first.write_text("first", encoding="utf-8")
    second.write_text("second", encoding="utf-8")
    started = threading.Event()
    release = threading.Event()

    def blocked_parse(*args, **kwargs):
        started.set()
        release.wait(timeout=5)
        from dbgpt_app.attachment import AttachmentParseService

        return AttachmentParseService(storage_dir=tmp_path / "cache").parse_files(
            *args, **kwargs
        )

    first_manager = AttachmentParseJobManager(
        storage_dir=tmp_path / "jobs",
        max_workers=1,
        max_concurrent_jobs_per_user=1,
        parse_batch=blocked_parse,
    )
    second_manager = AttachmentParseJobManager(
        storage_dir=tmp_path / "jobs",
        max_workers=1,
        max_concurrent_jobs_per_user=1,
    )
    first_manager.submit(
        user_id="alice", file_paths=[first], allowed_roots=[upload_dir]
    )
    assert started.wait(timeout=2)
    with pytest.raises(AttachmentJobLimitError):
        second_manager.submit(
            user_id="alice", file_paths=[second], allowed_roots=[upload_dir]
        )
    release.set()
    first_manager.shutdown()
    second_manager.shutdown()


def test_attachment_parse_job_retries_transient_worker_error(tmp_path, monkeypatch):
    import time

    from dbgpt_app.attachment import AttachmentParseService
    from dbgpt_app.attachment import storage
    from dbgpt_app.attachment.jobs import (
        AttachmentParseJobManager,
        TransientAttachmentJobError,
    )
    from dbgpt_app.attachment.storage import AttachmentChunkStore

    upload_dir = tmp_path / "uploads"
    upload_dir.mkdir()
    source = upload_dir / "retry.txt"
    source.write_text("retry succeeds", encoding="utf-8")
    monkeypatch.setattr(
        storage,
        "get_attachment_chunk_store",
        lambda: AttachmentChunkStore(storage_dir=tmp_path / "index"),
    )
    attempts = 0

    def flaky_parse(*args, **kwargs):
        nonlocal attempts
        attempts += 1
        if attempts == 1:
            raise TransientAttachmentJobError("temporary parser outage")
        kwargs.pop("cache_namespace", None)
        return AttachmentParseService(storage_dir=tmp_path / "cache").parse_files(
            *args, **kwargs
        )

    manager = AttachmentParseJobManager(
        storage_dir=tmp_path / "jobs",
        max_workers=1,
        max_retries=1,
        parse_batch=flaky_parse,
    )
    job = manager.submit(
        user_id="alice",
        file_paths=[source],
        attachment_ids=["att_retry"],
        allowed_roots=[upload_dir],
    )
    deadline = time.time() + 4
    while time.time() < deadline:
        current = manager.get(job.job_id, user_id="alice")
        if current.status in {"completed", "failed"}:
            break
        time.sleep(0.05)
    assert current.status == "completed"
    assert attempts == 2
    assert current.attempt == 2
    manager.shutdown()


def test_attachment_chunk_store_preserves_structure_and_user_scope(tmp_path):
    from dbgpt_app.attachment.models import AttachmentParseResult
    from dbgpt_app.attachment.storage import AttachmentChunkStore

    result = AttachmentParseResult(
        attachment_id="att_contract",
        filename="contract.md",
        format="markdown",
        mime_type="text/markdown",
        blocks=[
            {
                "block_id": "b_heading",
                "type": "heading",
                "text": "Payment terms",
                "location": {"line": 3},
            },
            {
                "block_id": "b_paragraph",
                "type": "paragraph",
                "text": "Payment is due within thirty working days after signing.",
                "location": {"line": 4},
            },
            {
                "block_id": "b_table",
                "type": "table",
                "text": "Item | Date\nInvoice | 2026-10-01",
                "columns": ["Item", "Date"],
                "rows": [["Invoice", "2026-10-01"]],
                "location": {"table": 1},
            },
        ],
    )
    store = AttachmentChunkStore(storage_dir=tmp_path / "index")

    assert store.replace_result("alice", result) == 3
    matches = store.search(
        "alice", "working days", attachment_ids=["att_contract"], top_k=3
    )
    assert matches[0]["block_id"] == "b_paragraph"
    assert matches[0]["location"] == {"line": 4}
    assert store.search("bob", "working days", attachment_ids=["att_contract"]) == []
    assert store.search(
        "alice", "invoice", attachment_ids=["att_contract"], include_tables=False
    ) == []
    assert AttachmentChunkStore(storage_dir=tmp_path / "index").search(
        "alice", "invoice", attachment_ids=["att_contract"]
    )[0]["type"] == "table"


def test_attachment_tools_parse_search_and_read_with_user_scope(tmp_path, monkeypatch):
    import json

    from dbgpt_app.attachment import AttachmentParseService
    from dbgpt_app.attachment.storage import AttachmentChunkStore
    from docx import Document
    from openpyxl import Workbook
    from dbgpt_app.openapi.api_v1 import python_upload_api
    from dbgpt_app.openapi.api_v1.tools import attachment_tools

    monkeypatch.setattr(
        python_upload_api.CFG,
        "SYSTEM_APP",
        SimpleNamespace(work_dir=str(tmp_path)),
    )
    monkeypatch.setattr(
        attachment_tools,
        "AttachmentParseService",
        lambda **kwargs: AttachmentParseService(
            storage_dir=tmp_path / "cache", **kwargs
        ),
    )
    index = AttachmentChunkStore(storage_dir=tmp_path / "index")
    monkeypatch.setattr(attachment_tools, "get_attachment_chunk_store", lambda: index)
    upload_dir = tmp_path / "python_uploads" / "alice"
    upload_dir.mkdir(parents=True)
    uploaded = upload_dir / "note.txt"
    uploaded.write_text("needle in line one\nsecond line", encoding="utf-8")
    document_path = upload_dir / "document.docx"
    document = Document()
    document.add_paragraph("First paragraph")
    document.add_paragraph("Target paragraph")
    table = document.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "Selected cell"
    document.save(document_path)
    spreadsheet_path = upload_dir / "sheet.xlsx"
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "销售"
    worksheet["A1"] = "区域"
    worksheet["B1"] = "销售额"
    worksheet["A2"] = "华东"
    worksheet["B2"] = "128万元"
    worksheet["A3"] = "华南"
    worksheet["B3"] = "96万元"
    workbook.save(spreadsheet_path)
    react_state = {
        "file_paths": [str(uploaded), str(document_path), str(spreadsheet_path)],
        "user_id": "alice",
        "session_state": {
            "attachments": [
                {
                    "attachment_id": "att_note",
                    "absolute_path": str(uploaded),
                    "filename": uploaded.name,
                },
                {
                    "attachment_id": "att_docx",
                    "absolute_path": str(document_path),
                    "filename": document_path.name,
                },
                {
                    "attachment_id": "att_xlsx",
                    "absolute_path": str(spreadsheet_path),
                    "filename": spreadsheet_path.name,
                },
            ]
        },
    }
    list_tool, parse_tool, search_tool, read_tool, export_tool = (
        attachment_tools.make_attachment_tools(react_state)
    )

    listing = json.loads(list_tool())["chunks"][0]["content"]
    parsed = json.loads(parse_tool())["chunks"][0]["content"][0]
    matches = json.loads(search_tool(query="needle"))["chunks"][0]["content"]
    region = json.loads(read_tool(line_start=2, line_end=2))["chunks"][0]["content"]
    paragraph_region = json.loads(read_tool(file_path=str(document_path), paragraph=2))[
        "chunks"
    ][0]["content"]
    table_region = json.loads(
        read_tool(file_path=str(document_path), table=1, block_type="table")
    )["chunks"][0]["content"]
    exported = json.loads(
        export_tool(file_path=str(document_path), format="markdown")
    )["chunks"][0]["content"]
    id_region = json.loads(
        read_tool(attachment_id="att_docx", paragraph=2)
    )["chunks"][0]["content"]
    spreadsheet_region = json.loads(
        read_tool(
            file_path=str(spreadsheet_path),
            sheet="销售",
            row_start=2,
            row_end=2,
        )
    )["chunks"][0]["content"]

    assert listing[0]["filename"] == "note.txt"
    assert "path" not in listing[0]
    assert str(uploaded) not in json.dumps(listing)
    assert parsed["status"] == "completed"
    assert "blocks" not in parsed
    assert matches[0]["text"] == "needle in line one"
    assert matches[0]["attachment_id"] == "att_note"
    assert region[0]["location"]["line"] == 2
    assert region[0]["text"] == "second line"
    assert paragraph_region[0]["text"] == "Target paragraph"
    assert table_region[0]["text"] == "Selected cell"
    assert "Target paragraph" in exported[0]["content"]
    assert id_region[0]["text"] == "Target paragraph"
    assert spreadsheet_region[0]["location"]["cell_range"] == "销售!A2:B2"
    assert spreadsheet_region[0]["cells"][1]["value"] == "128万元"
