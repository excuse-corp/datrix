import os
import time
import zipfile
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path

import pytest

from dbgpt_app.attachment import AttachmentParseService, ParseOptions
from dbgpt_app.attachment.registry import UnsupportedAttachmentFormat, detect_format
from dbgpt_app.attachment.service import AttachmentParsingCancelled


def make_service(tmp_path):
    return AttachmentParseService(storage_dir=tmp_path / "cache")


def test_txt_utf8_and_gb18030(tmp_path):
    service = make_service(tmp_path)
    utf8 = tmp_path / "utf8.txt"
    utf8.write_text("第一行\n第二行", encoding="utf-8")
    result = service.parse_file(utf8)
    assert result.status == "completed"
    assert result.format == "txt"
    assert "第一行" in result.plain_text

    gb = tmp_path / "legacy.txt"
    gb.write_bytes("中文文本".encode("gb18030"))
    result = service.parse_file(gb)
    assert result.status == "completed"
    assert "中文文本" in result.plain_text
    assert result.metadata["encoding"].lower().replace("-", "") in {
        "gb18030",
        "gb2312",
        "gbk",
    }

    classic_mac = tmp_path / "classic-mac.txt"
    classic_mac.write_bytes(b"first\rsecond")
    result = service.parse_file(classic_mac, persist=False)
    assert result.metadata["line_count"] == 2
    assert [item["text"] for item in result.blocks] == ["first", "second"]


def test_markdown_headings_code_and_table(tmp_path):
    path = tmp_path / "README.MARKDOWN"
    path.write_text(
        "# Title\n\n| Name | Count |\n| --- | ---: |\n| A | 2 |\n\n"
        "```python\nprint(1)\n```",
        encoding="utf-8",
    )
    result = make_service(tmp_path).parse_file(path)
    assert result.format == "md"
    assert {item["type"] for item in result.blocks} >= {"heading", "table", "code"}
    table = next(item for item in result.blocks if item["type"] == "table")
    assert table["columns"] == ["Name", "Count"]
    assert table["rows"] == [["A", "2"]]


def test_html_removes_active_content_and_extracts_table(tmp_path):
    path = tmp_path / "page.HTM"
    path.write_text(
        "<html><head><title>Page</title><script>alert(1)</script></head>"
        "<body><h1>Heading</h1><p>Hello</p><table><tr><th>A</th><th>B</th></tr>"
        "<tr><td>1</td><td>2</td></tr></table></body></html>",
        encoding="utf-8",
    )
    result = make_service(tmp_path).parse_file(path)
    assert result.format == "html"
    assert "alert" not in result.plain_text
    assert "Hello" in result.plain_text
    table = next(item for item in result.blocks if item["type"] == "table")
    assert table["rows"] == [["1", "2"]]


def test_docx_paragraph_heading_and_table(tmp_path):
    from docx import Document

    path = tmp_path / "document.docx"
    document = Document()
    document.add_heading("Heading", level=1)
    document.add_paragraph("Body paragraph")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "A"
    table.cell(0, 1).text = "B"
    table.cell(1, 0).text = "1"
    table.cell(1, 1).text = "{{销售额}}"
    document.save(path)

    result = make_service(tmp_path).parse_file(path)
    assert result.format == "docx"
    assert any(item["type"] == "heading" for item in result.blocks)
    table_block = next(item for item in result.blocks if item["type"] == "table")
    assert table_block["location"] == {
        "part": "body",
        "table": 1,
        "row_count": 2,
        "column_count": 2,
    }
    assert table_block["cells"][1][1]["text"] == "{{销售额}}"
    assert table_block["cells"][1][1]["placeholders"] == ["销售额"]
    assert result.metadata["placeholders"] == ["销售额"]
    assert "Body paragraph" in result.plain_text


def test_xlsx_sheet_cells_formulas_and_placeholders(tmp_path):
    from openpyxl import Workbook
    from openpyxl.styles import PatternFill

    path = tmp_path / "template.xlsx"
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "销售"
    worksheet["A1"] = "区域"
    worksheet["B1"] = "销售额"
    worksheet["A2"] = "华东"
    worksheet["B2"].fill = PatternFill(fill_type="solid", fgColor="FFFF00")
    worksheet["D1"] = "{{销售额}}"
    worksheet["F1"] = "=SUM(B2:B3)"
    worksheet.merge_cells("D2:E2")
    worksheet["D2"] = "合计"
    workbook.create_sheet("说明")["A1"] = "填写说明"
    workbook.save(path)

    result = make_service(tmp_path).parse_file(path, persist=False)

    assert result.status == "completed"
    assert result.format == "xlsx"
    assert result.parser == "excel"
    assert result.metadata["sheet_names"] == ["销售", "说明"]
    assert result.metadata["placeholders"] == ["销售额"]
    assert result.metadata["formula_count"] == 1

    rows = [
        item
        for item in result.blocks
        if item.get("type") == "table" and item.get("sheet") == "销售"
    ]
    row_two = next(item for item in rows if item["row_number"] == 2)
    blank_cell = next(cell for cell in row_two["cells"] if cell["coordinate"] == "B2")
    assert blank_cell["empty"] is True
    assert row_two["location"]["cell_range"] == "销售!A2:D2"

    row_one = next(item for item in rows if item["row_number"] == 1)
    formula_cell = next(cell for cell in row_one["cells"] if cell["coordinate"] == "F1")
    assert formula_cell["is_formula"] is True
    assert formula_cell["formula"] == "=SUM(B2:B3)"

    worksheet_block = next(
        item
        for item in result.blocks
        if item.get("type") == "worksheet" and item.get("sheet") == "销售"
    )
    assert "销售!D2:E2" in worksheet_block["merged_ranges"]


def test_pptx_slide_and_table(tmp_path):
    from pptx import Presentation
    from pptx.util import Inches

    path = tmp_path / "slides.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    textbox = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
    textbox.text = "Slide text"
    table_shape = slide.shapes.add_table(
        2, 2, Inches(1), Inches(2), Inches(4), Inches(2)
    )
    table_shape.table.cell(0, 0).text = "A"
    table_shape.table.cell(0, 1).text = "B"
    table_shape.table.cell(1, 0).text = "1"
    table_shape.table.cell(1, 1).text = "2"
    slide.notes_slide.notes_text_frame.text = "Speaker note"
    presentation.save(path)

    result = make_service(tmp_path).parse_file(path)
    assert result.format == "pptx"
    assert result.metadata["slide_count"] == 1
    assert any(item["type"] == "slide" for item in result.blocks)
    assert any(item["type"] == "table" for item in result.blocks)
    assert any(
        item["type"] == "speaker_notes" and item["text"] == "Speaker note"
        for item in result.blocks
    )


def test_pdf_textless_page_requires_ocr(tmp_path):
    from pypdf import PdfWriter

    path = tmp_path / "scan.PDF"
    writer = PdfWriter()
    writer.add_blank_page(width=612, height=792)
    with path.open("wb") as handle:
        writer.write(handle)

    result = make_service(tmp_path).parse_file(path)
    assert result.format == "pdf"
    assert result.metadata["page_count"] == 1
    assert result.status == "needs_ocr"


def test_pdf_text_extraction_and_doc_adapter(tmp_path):
    repository_root = Path(__file__).resolve().parents[5]
    pdf_path = repository_root / "examples/agents/example_files/Nuclear_power.pdf"
    doc_path = (
        repository_root
        / "packages/dbgpt-ext/src/dbgpt_ext/rag/knowledge/tests/data/test_mock.doc"
    )
    service = make_service(tmp_path)
    pdf_result = service.parse_file(
        pdf_path,
        options=ParseOptions(extract_tables=False),
        persist=False,
    )
    doc_result = service.parse_file(doc_path, persist=False)
    assert pdf_result.status == "completed"
    assert pdf_result.metadata["page_count"] > 0
    assert pdf_result.plain_text.strip()
    assert doc_result.format == "doc"
    assert doc_result.blocks


def test_unsupported_corrupt_and_uppercase_detection(tmp_path):
    unknown = tmp_path / "data.bin"
    unknown.write_bytes(b"not a supported document")
    with pytest.raises(UnsupportedAttachmentFormat):
        detect_format(str(unknown))

    corrupt = tmp_path / "bad.docx"
    corrupt.write_bytes(b"PK\x03\x04broken")
    result = make_service(tmp_path).parse_file(corrupt)
    assert result.status == "failed"
    assert result.error_code in {"parse_error", "unsupported_format"}

    pdf = tmp_path / "UPPER.PDF"
    pdf.write_bytes(b"%PDF-1.7\n")
    assert detect_format(str(pdf)) == "pdf"

    from openpyxl import Workbook

    workbook_path = tmp_path / "UPPER.XLSX"
    workbook = Workbook()
    workbook.active["A1"] = "Excel"
    workbook.save(workbook_path)
    assert detect_format(str(workbook_path)) == "xlsx"

    mislabeled = tmp_path / "document.txt"
    with zipfile.ZipFile(mislabeled, "w") as archive:
        archive.writestr("word/document.xml", "<document />")
    with pytest.raises(UnsupportedAttachmentFormat, match="does not match"):
        detect_format(str(mislabeled))

    image_as_text = tmp_path / "picture.txt"
    image_as_text.write_bytes(b"\x89PNG\r\n\x1a\ncontent")
    with pytest.raises(UnsupportedAttachmentFormat, match="Binary files"):
        detect_format(str(image_as_text))


def test_path_traversal_zip_limits_and_batch_isolation(tmp_path, monkeypatch):
    allowed = tmp_path / "uploads"
    allowed.mkdir()
    good = allowed / "good.txt"
    good.write_text("safe", encoding="utf-8")
    outside = tmp_path / "outside.txt"
    outside.write_text("outside", encoding="utf-8")
    service = make_service(tmp_path)

    result = service.parse_files(
        [str(good), str(allowed / ".." / "outside.txt")],
        allowed_roots=[allowed],
    )
    assert [item.status for item in result] == ["completed", "failed"]
    assert result[1].error_code == "attachment_rejected"

    from dbgpt_app.attachment import registry

    monkeypatch.setattr(registry, "MAX_ZIP_UNCOMPRESSED_BYTES", 5)
    zip_path = tmp_path / "many.docx"
    with zipfile.ZipFile(zip_path, "w") as archive:
        for index in range(5001):
            archive.writestr(f"entry-{index}", b"x")
    with pytest.raises(UnsupportedAttachmentFormat):
        detect_format(str(zip_path))

    zip_bomb = tmp_path / "expanded.docx"
    with zipfile.ZipFile(zip_bomb, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", b"x" * 6)
    with pytest.raises(UnsupportedAttachmentFormat, match="uncompressed size"):
        detect_format(str(zip_bomb))


def test_text_and_block_limits_and_cache(tmp_path):
    path = tmp_path / "many.txt"
    path.write_text("a\n" * 20, encoding="utf-8")
    service = make_service(tmp_path)
    block_options = ParseOptions(max_text_chars=1000, max_blocks=3)
    block_limited = service.parse_file(path, options=block_options, persist=False)
    assert len(block_limited.blocks) == 3
    assert any("block limit" in warning.lower() for warning in block_limited.warnings)

    text_options = ParseOptions(max_text_chars=5, max_blocks=50)
    first = service.parse_file(path, options=text_options)
    second = service.parse_file(path, options=text_options)
    assert len(first.plain_text) <= 5
    assert any("text limit" in warning.lower() for warning in first.warnings)
    assert second.attachment_id == first.attachment_id
    assert second.blocks == first.blocks
    assert first.blocks[0]["block_id"] == "b_00001"
    assert list((tmp_path / "cache").glob("*.json"))


def test_file_size_limit_returns_failed_batch_item(tmp_path):
    path = tmp_path / "too-large.txt"
    path.write_text("larger than limit", encoding="utf-8")
    service = AttachmentParseService(
        storage_dir=tmp_path / "cache",
        max_file_size_bytes=2,
    )
    result = service.parse_files([path])[0]
    assert result.status == "failed"
    assert result.error_code == "attachment_rejected"


def test_document_parser_timeout_returns_failed_result(tmp_path):
    path = tmp_path / "timeout.pdf"
    path.write_bytes(b"%PDF-1.7\n")
    service = AttachmentParseService(
        storage_dir=tmp_path / "cache",
        parser_timeout_seconds=1,
    )
    service._worker_bootstrap = (
        "import json, time; json.load(__import__('sys').stdin); time.sleep(30)"
    )
    result = service.parse_file(path, persist=False)
    assert result.status == "failed"
    assert result.error_code == "parser_timeout"


def test_cancelling_parse_kills_worker_process(tmp_path, monkeypatch):
    pid_path = tmp_path / "worker.pid"
    child_pid_path = tmp_path / "worker-child.pid"
    monkeypatch.setenv("DATRIX_TEST_ATTACHMENT_WORKER_PID", str(pid_path))
    monkeypatch.setenv("DATRIX_TEST_ATTACHMENT_CHILD_PID", str(child_pid_path))
    path = tmp_path / "cancel.pdf"
    path.write_bytes(b"%PDF-1.7\n")
    service = AttachmentParseService(
        storage_dir=tmp_path / "cache",
        parser_timeout_seconds=30,
    )
    service._worker_bootstrap = (
        "import json, os, subprocess, sys, time; json.load(sys.stdin); "
        "open(os.environ['DATRIX_TEST_ATTACHMENT_WORKER_PID'], 'w').write(str(os.getpid())); "
        "child = subprocess.Popen([sys.executable, '-c', \"import os, time; "
        "open(os.environ['DATRIX_TEST_ATTACHMENT_CHILD_PID'], 'w').write(str(os.getpid())); "
        'time.sleep(30)"]); '
        "time.sleep(30)"
    )

    with ThreadPoolExecutor(max_workers=1) as executor:
        future = executor.submit(service.parse_file, path, persist=False)
        deadline = time.monotonic() + 5
        while not pid_path.exists() and time.monotonic() < deadline:
            time.sleep(0.01)
        assert pid_path.exists(), "parser worker did not start"
        while not child_pid_path.exists() and time.monotonic() < deadline:
            time.sleep(0.01)
        assert child_pid_path.exists(), "parser worker child did not start"
        worker_pid = int(pid_path.read_text(encoding="utf-8"))
        child_pid = int(child_pid_path.read_text(encoding="utf-8"))
        service.cancel_active_parsers()
        with pytest.raises(AttachmentParsingCancelled):
            future.result(timeout=5)

    with pytest.raises(ProcessLookupError):
        os.kill(worker_pid, 0)
    with pytest.raises(ProcessLookupError):
        os.kill(child_pid, 0)
    assert not service._active_processes


def test_cancelled_service_does_not_start_new_worker(tmp_path):
    path = tmp_path / "cancel-before-start.pdf"
    path.write_bytes(b"%PDF-1.7\n")
    service = AttachmentParseService(storage_dir=tmp_path / "cache")
    service.cancel_active_parsers()

    with pytest.raises(AttachmentParsingCancelled):
        service.parse_file(path, persist=False)
