"""Excel parser with worksheet, row, and cell-level locations."""

from __future__ import annotations

from datetime import date, datetime, time
from pathlib import Path
from typing import Any, Dict, Iterable, List

from dbgpt_app.attachment.models import AttachmentParseResult, ParseOptions

from .base import block, finalize_result


MAX_SHEETS = 100
MAX_ROWS_PER_SHEET = 100_000
MAX_COLUMNS_PER_SHEET = 512
MAX_CELLS_PER_WORKBOOK = 250_000
MAX_SCANNED_CELLS = 1_000_000
MAX_MERGED_RANGES = 2_000


def _json_value(value: Any) -> Any:
    """Convert spreadsheet values to values safe for JSON serialization."""

    if isinstance(value, datetime):
        return value.isoformat(sep=" ")
    if isinstance(value, (date, time)):
        return value.isoformat()
    if isinstance(value, bytes):
        return value.decode("utf-8", errors="replace")
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    return str(value)


def _display_value(value: Any) -> str:
    normalized = _json_value(value)
    if normalized is None:
        return ""
    return str(normalized)


def _column_letter(column: int) -> str:
    letters = ""
    current = max(1, int(column))
    while current:
        current, remainder = divmod(current - 1, 26)
        letters = chr(65 + remainder) + letters
    return letters


def _cell_range(sheet: str, row: int, first_column: int, last_column: int) -> str:
    start = f"{_column_letter(first_column)}{row}"
    end = f"{_column_letter(last_column)}{row}"
    return f"{sheet}!{start}:{end}"


def _placeholder_tokens(value: Any) -> List[str]:
    import re

    text = _display_value(value)
    tokens: List[str] = []
    for pattern in (r"\{\{\s*([^{}]+?)\s*\}\}", r"<<\s*([^<>]+?)\s*>>"):
        for match in re.finditer(pattern, text):
            token = match.group(1).strip()
            if token and token not in tokens:
                tokens.append(token)
    return tokens


def _merged_range_strings(ranges: Iterable[Any], sheet_name: str) -> List[str]:
    values: List[str] = []
    for item in ranges:
        if len(values) >= MAX_MERGED_RANGES:
            break
        values.append(f"{sheet_name}!{str(item)}")
    return values


def _cell_payload(
    *,
    coordinate: str,
    row: int,
    column: int,
    value: Any,
    data_type: str | None = None,
    number_format: str | None = None,
    style_id: int | None = None,
    is_formula: bool = False,
) -> Dict[str, Any]:
    normalized = _json_value(value)
    payload: Dict[str, Any] = {
        "coordinate": coordinate,
        "row": row,
        "column": column,
        "value": normalized,
        "display_value": _display_value(value),
        "empty": normalized is None or normalized == "",
    }
    if data_type:
        payload["data_type"] = data_type
    if number_format and number_format != "General":
        payload["number_format"] = number_format
    if style_id:
        payload["style_id"] = style_id
    if is_formula:
        payload["is_formula"] = True
        payload["formula"] = _display_value(value)
    placeholders = _placeholder_tokens(value)
    if placeholders:
        payload["placeholders"] = placeholders
    return payload


def _row_block(
    *,
    sheet_name: str,
    sheet_index: int,
    row_number: int,
    cells: List[Dict[str, Any]],
) -> Dict[str, Any]:
    first_column = min(int(cell["column"]) for cell in cells)
    last_column = max(int(cell["column"]) for cell in cells)
    values = [cell.get("value") for cell in cells]
    cell_text = " | ".join(
        f"{cell['coordinate']}={cell.get('display_value', '')}" for cell in cells
    )
    placeholders: List[str] = []
    for cell in cells:
        for token in cell.get("placeholders") or []:
            if token not in placeholders:
                placeholders.append(token)
    return block(
        "table",
        f"{sheet_name} row {row_number}: {cell_text}",
        location={
            "sheet": sheet_name,
            "sheet_index": sheet_index,
            "row": row_number,
            "cell_range": _cell_range(
                sheet_name, row_number, first_column, last_column
            ),
        },
        sheet=sheet_name,
        sheet_index=sheet_index,
        row_number=row_number,
        columns=[cell["coordinate"] for cell in cells],
        rows=[values],
        cells=cells,
        placeholders=placeholders,
    )


class ExcelParser:
    """Parse XLSX/XLS files into searchable rows with cell coordinates."""

    name = "excel"

    def parse(
        self,
        path: Path,
        attachment_id: str,
        mime_type: str | None,
        options: ParseOptions,
    ) -> AttachmentParseResult:
        extension = path.suffix.lower()
        if extension == ".xlsx":
            return self._parse_xlsx(path, attachment_id, mime_type, options)
        if extension == ".xls":
            return self._parse_xls(path, attachment_id, mime_type, options)
        raise ValueError(f"Unsupported Excel extension: {extension or 'unknown'}")

    def _parse_xlsx(
        self,
        path: Path,
        attachment_id: str,
        mime_type: str | None,
        options: ParseOptions,
    ) -> AttachmentParseResult:
        from openpyxl import load_workbook

        workbook = load_workbook(
            filename=str(path),
            read_only=False,
            data_only=False,
            keep_links=False,
        )
        try:
            return self._parse_openpyxl_workbook(
                workbook, path, attachment_id, mime_type, options, "xlsx"
            )
        finally:
            workbook.close()

    def _parse_openpyxl_workbook(
        self,
        workbook: Any,
        path: Path,
        attachment_id: str,
        mime_type: str | None,
        options: ParseOptions,
        format_name: str,
    ) -> AttachmentParseResult:
        blocks: List[Dict[str, Any]] = []
        warnings: List[str] = []
        sheet_summaries: List[Dict[str, Any]] = []
        placeholder_tokens: List[str] = []
        formula_count = 0
        parsed_cells = 0
        scanned_cells = 0
        sheet_count = min(len(workbook.worksheets), MAX_SHEETS)
        if len(workbook.worksheets) > MAX_SHEETS:
            warnings.append(f"Sheet limit reached: {MAX_SHEETS}")

        for sheet_index, worksheet in enumerate(workbook.worksheets[:sheet_count], start=1):
            if parsed_cells >= MAX_CELLS_PER_WORKBOOK:
                warnings.append(f"Cell limit reached: {MAX_CELLS_PER_WORKBOOK}")
                break

            sheet_name = str(worksheet.title)
            max_row = min(int(worksheet.max_row or 0), MAX_ROWS_PER_SHEET)
            max_column = min(int(worksheet.max_column or 0), MAX_COLUMNS_PER_SHEET)
            if (worksheet.max_row or 0) > MAX_ROWS_PER_SHEET:
                warnings.append(
                    f"Row limit reached for sheet {sheet_name}: {MAX_ROWS_PER_SHEET}"
                )
            if (worksheet.max_column or 0) > MAX_COLUMNS_PER_SHEET:
                warnings.append(
                    f"Column limit reached for sheet {sheet_name}: "
                    f"{MAX_COLUMNS_PER_SHEET}"
                )

            try:
                dimension = worksheet.calculate_dimension()
            except Exception:
                dimension = None
            merged_ranges = _merged_range_strings(
                list(worksheet.merged_cells.ranges), sheet_name
            )
            sheet_placeholders: List[str] = []
            sheet_summaries.append(
                {
                    "name": sheet_name,
                    "index": sheet_index,
                    "state": getattr(worksheet, "sheet_state", "visible"),
                    "dimensions": dimension,
                    "rows": max_row,
                    "columns": max_column,
                    "merged_ranges": merged_ranges,
                }
            )
            metadata_text = (
                f"Worksheet {sheet_name}: {max_row} rows x {max_column} columns"
            )
            blocks.append(
                block(
                    "worksheet",
                    metadata_text,
                    location={"sheet": sheet_name, "sheet_index": sheet_index},
                    sheet=sheet_name,
                    sheet_index=sheet_index,
                    dimensions=dimension,
                    row_count=max_row,
                    column_count=max_column,
                    merged_ranges=merged_ranges,
                )
            )

            if not options.extract_tables or not max_row or not max_column:
                continue

            for row_number, row in enumerate(
                worksheet.iter_rows(
                    min_row=1,
                    max_row=max_row,
                    min_col=1,
                    max_col=max_column,
                ),
                start=1,
            ):
                if parsed_cells >= MAX_CELLS_PER_WORKBOOK:
                    warnings.append(f"Cell limit reached: {MAX_CELLS_PER_WORKBOOK}")
                    break

                active_cells = []
                for cell in row:
                    scanned_cells += 1
                    if scanned_cells > MAX_SCANNED_CELLS:
                        warnings.append(f"Scan limit reached: {MAX_SCANNED_CELLS}")
                        break
                    value = cell.value
                    has_style = bool(getattr(cell, "has_style", False))
                    if value is None and not has_style:
                        continue
                    is_formula = getattr(cell, "data_type", None) == "f"
                    if is_formula:
                        formula_count += 1
                    active_cells.append(
                        _cell_payload(
                            coordinate=str(cell.coordinate),
                            row=row_number,
                            column=int(cell.column),
                            value=value,
                            data_type=getattr(cell, "data_type", None),
                            number_format=getattr(cell, "number_format", None),
                            style_id=getattr(cell, "style_id", None),
                            is_formula=is_formula,
                        )
                    )
                    parsed_cells += 1
                    for token in _placeholder_tokens(value):
                        if token not in sheet_placeholders:
                            sheet_placeholders.append(token)
                        if token not in placeholder_tokens:
                            placeholder_tokens.append(token)
                if active_cells:
                    blocks.append(
                        _row_block(
                            sheet_name=sheet_name,
                            sheet_index=sheet_index,
                            row_number=row_number,
                            cells=active_cells,
                        )
                    )
                if scanned_cells >= MAX_SCANNED_CELLS:
                    break

            for item in blocks:
                if item.get("sheet") == sheet_name and item.get("type") == "worksheet":
                    item["placeholders"] = sheet_placeholders
                    break

        metadata = {
            "sheet_count": len(sheet_summaries),
            "sheet_names": [item["name"] for item in sheet_summaries],
            "sheets": sheet_summaries,
            "formula_count": formula_count,
            "placeholder_count": len(placeholder_tokens),
            "placeholders": placeholder_tokens,
            "parsed_cell_count": parsed_cells,
            "format": format_name,
        }
        return finalize_result(
            attachment_id=attachment_id,
            path=path,
            format_name=format_name,
            mime_type=mime_type,
            parser=self.name,
            blocks=blocks,
            metadata=metadata,
            warnings=warnings,
            options=options,
        )

    def _parse_xls(
        self,
        path: Path,
        attachment_id: str,
        mime_type: str | None,
        options: ParseOptions,
    ) -> AttachmentParseResult:
        import xlrd

        workbook = xlrd.open_workbook(str(path), on_demand=False)
        blocks: List[Dict[str, Any]] = []
        warnings: List[str] = []
        sheet_summaries: List[Dict[str, Any]] = []
        placeholder_tokens: List[str] = []
        parsed_cells = 0
        scanned_cells = 0
        formula_count = 0
        sheet_count = min(workbook.nsheets, MAX_SHEETS)
        if workbook.nsheets > MAX_SHEETS:
            warnings.append(f"Sheet limit reached: {MAX_SHEETS}")

        for sheet_index in range(sheet_count):
            if parsed_cells >= MAX_CELLS_PER_WORKBOOK:
                warnings.append(f"Cell limit reached: {MAX_CELLS_PER_WORKBOOK}")
                break
            worksheet = workbook.sheet_by_index(sheet_index)
            sheet_name = str(worksheet.name)
            row_count = min(int(worksheet.nrows), MAX_ROWS_PER_SHEET)
            column_count = min(int(worksheet.ncols), MAX_COLUMNS_PER_SHEET)
            merged_ranges = []
            for row_start, row_end, col_start, col_end in worksheet.merged_cells[
                :MAX_MERGED_RANGES
            ]:
                merged_ranges.append(
                    f"{sheet_name}!{_column_letter(col_start + 1)}{row_start + 1}:"
                    f"{_column_letter(col_end)}{row_end}"
                )
            sheet_summaries.append(
                {
                    "name": sheet_name,
                    "index": sheet_index + 1,
                    "rows": row_count,
                    "columns": column_count,
                    "merged_ranges": merged_ranges,
                }
            )
            blocks.append(
                block(
                    "worksheet",
                    f"Worksheet {sheet_name}: {row_count} rows x {column_count} columns",
                    location={"sheet": sheet_name, "sheet_index": sheet_index + 1},
                    sheet=sheet_name,
                    sheet_index=sheet_index + 1,
                    row_count=row_count,
                    column_count=column_count,
                    merged_ranges=merged_ranges,
                )
            )
            if not options.extract_tables:
                continue

            for row_number in range(row_count):
                if parsed_cells >= MAX_CELLS_PER_WORKBOOK:
                    warnings.append(f"Cell limit reached: {MAX_CELLS_PER_WORKBOOK}")
                    break
                row_cells = worksheet.row(row_number)
                active_columns = []
                for column_number in range(column_count):
                    scanned_cells += 1
                    if scanned_cells > MAX_SCANNED_CELLS:
                        warnings.append(f"Scan limit reached: {MAX_SCANNED_CELLS}")
                        break
                    cell = row_cells[column_number]
                    value = cell.value
                    styled_blank = False
                    try:
                        styled_blank = cell.ctype == xlrd.XL_CELL_BLANK and bool(
                            worksheet.cell_xf_index(row_number, column_number)
                        )
                    except Exception:
                        pass
                    if value in (None, "") and not styled_blank:
                        continue
                    if cell.ctype == xlrd.XL_CELL_FORMULA:
                        formula_count += 1
                    active_columns.append(
                        _cell_payload(
                            coordinate=f"{_column_letter(column_number + 1)}{row_number + 1}",
                            row=row_number + 1,
                            column=column_number + 1,
                            value=value,
                            data_type=str(cell.ctype),
                            is_formula=cell.ctype == xlrd.XL_CELL_FORMULA,
                        )
                    )
                    parsed_cells += 1
                    for token in _placeholder_tokens(value):
                        if token not in placeholder_tokens:
                            placeholder_tokens.append(token)
                if active_columns:
                    blocks.append(
                        _row_block(
                            sheet_name=sheet_name,
                            sheet_index=sheet_index + 1,
                            row_number=row_number + 1,
                            cells=active_columns,
                        )
                    )
                if scanned_cells >= MAX_SCANNED_CELLS:
                    break

        metadata = {
            "sheet_count": len(sheet_summaries),
            "sheet_names": [item["name"] for item in sheet_summaries],
            "sheets": sheet_summaries,
            "formula_count": formula_count,
            "placeholder_count": len(placeholder_tokens),
            "placeholders": placeholder_tokens,
            "parsed_cell_count": parsed_cells,
            "format": "xls",
        }
        return finalize_result(
            attachment_id=attachment_id,
            path=path,
            format_name="xls",
            mime_type=mime_type,
            parser=self.name,
            blocks=blocks,
            metadata=metadata,
            warnings=warnings,
            options=options,
        )
