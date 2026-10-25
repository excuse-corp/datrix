"""PPTX parser preserving slide boundaries and tables."""

from pathlib import Path
from typing import Any, Dict, List

from dbgpt_app.attachment.models import AttachmentParseResult, ParseOptions

from .base import block, finalize_result


class PptxParser:
    name = "pptx"

    def parse(
        self,
        path: Path,
        attachment_id: str,
        mime_type: str | None,
        options: ParseOptions,
    ) -> AttachmentParseResult:
        from pptx import Presentation

        presentation = Presentation(str(path))
        blocks: List[Dict[str, Any]] = []
        warnings: List[str] = []
        for slide_number, slide in enumerate(presentation.slides, start=1):
            if len(blocks) >= options.max_blocks:
                warnings.append(f"Block limit reached: {options.max_blocks}")
                break
            slide_text: List[str] = []
            table_index = 0
            for shape in slide.shapes:
                if len(blocks) >= options.max_blocks:
                    warnings.append(f"Block limit reached: {options.max_blocks}")
                    break
                if getattr(shape, "has_table", False) and options.extract_tables:
                    table_index += 1
                    rows = [
                        [cell.text.strip() for cell in row.cells]
                        for row in shape.table.rows
                    ]
                    if rows:
                        blocks.append(
                            block(
                                "table",
                                "\n".join(" | ".join(row) for row in rows),
                                location={
                                    "slide": slide_number,
                                    "table": table_index,
                                },
                                columns=rows[0],
                                rows=rows[1:],
                            )
                        )
                elif getattr(shape, "has_text_frame", False):
                    text = shape.text.strip()
                    if text:
                        slide_text.append(text)
            if slide_text:
                blocks.append(
                    block(
                        "slide",
                        "\n".join(slide_text),
                        location={"slide": slide_number},
                    )
                )
            try:
                notes_text = slide.notes_slide.notes_text_frame.text.strip()
            except (AttributeError, ValueError):
                notes_text = ""
            if notes_text:
                blocks.append(
                    block(
                        "speaker_notes",
                        notes_text,
                        location={"slide": slide_number},
                    )
                )

        return finalize_result(
            attachment_id=attachment_id,
            path=path,
            format_name="pptx",
            mime_type=mime_type,
            parser=self.name,
            blocks=blocks,
            metadata={"slide_count": len(presentation.slides)},
            warnings=warnings,
            options=options,
        )
