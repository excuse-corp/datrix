"""DOCX parser preserving document order and editable table locations."""

from __future__ import annotations

import re
from pathlib import Path
from typing import Any, Dict, Iterable, List

from dbgpt_app.attachment.models import AttachmentParseResult, ParseOptions

from .base import block, finalize_result


_PLACEHOLDER_PATTERNS = (
    re.compile(r"\{\{\s*([^{}]+?)\s*\}\}"),
    re.compile(r"<<\s*([^<>]+?)\s*>>"),
)


def _placeholders(text: str) -> List[str]:
    values: List[str] = []
    for pattern in _PLACEHOLDER_PATTERNS:
        for match in pattern.finditer(text or ""):
            value = match.group(1).strip()
            if value and value not in values:
                values.append(value)
    return values


def _cell_payload(cell: Any, row_number: int, column_number: int) -> Dict[str, Any]:
    text = cell.text.strip()
    payload: Dict[str, Any] = {
        "row": row_number,
        "column": column_number,
        "text": text,
        "empty": not bool(text),
    }
    placeholders = _placeholders(text)
    if placeholders:
        payload["placeholders"] = placeholders
    return payload


def _table_payload(table: Any, table_index: int, part: str = "body") -> Dict[str, Any]:
    rows: List[List[str]] = []
    cells: List[List[Dict[str, Any]]] = []
    placeholders: List[str] = []
    max_columns = 0
    for row_number, row in enumerate(table.rows, start=1):
        row_values: List[str] = []
        row_cells: List[Dict[str, Any]] = []
        for column_number, cell in enumerate(row.cells, start=1):
            payload = _cell_payload(cell, row_number, column_number)
            row_cells.append(payload)
            row_values.append(payload["text"])
            for token in payload.get("placeholders") or []:
                if token not in placeholders:
                    placeholders.append(token)
        rows.append(row_values)
        cells.append(row_cells)
        max_columns = max(max_columns, len(row_values))

    if not rows:
        return {}
    return {
        "table_index": table_index,
        "part": part,
        "row_count": len(rows),
        "column_count": max_columns,
        "rows": rows,
        "cells": cells,
        "placeholders": placeholders,
    }


def _paragraph_payload(paragraph: Any, *, location: Dict[str, Any]) -> Dict[str, Any]:
    text = paragraph.text.strip()
    style_name = paragraph.style.name if paragraph.style else ""
    payload: Dict[str, Any] = {
        "text": text,
        "style": style_name,
        "location": location,
    }
    placeholders = _placeholders(text)
    if placeholders:
        payload["placeholders"] = placeholders
    return payload


class DocxParser:
    name = "docx"

    def parse(
        self,
        path: Path,
        attachment_id: str,
        mime_type: str | None,
        options: ParseOptions,
    ) -> AttachmentParseResult:
        from docx import Document
        from docx.table import Table
        from docx.text.paragraph import Paragraph

        document = Document(str(path))
        blocks: List[Dict[str, Any]] = []
        warnings: List[str] = []
        paragraph_index = 0
        table_index = 0
        all_placeholders: List[str] = []

        for child in document.element.body.iterchildren():
            if len(blocks) >= options.max_blocks:
                warnings.append(f"Block limit reached: {options.max_blocks}")
                break
            if child.tag.endswith("}p"):
                paragraph = Paragraph(child, document)
                text = paragraph.text.strip()
                if not text:
                    continue
                paragraph_index += 1
                style_name = paragraph.style.name if paragraph.style else ""
                paragraph_data = _paragraph_payload(
                    paragraph,
                    location={
                        "part": "body",
                        "paragraph": paragraph_index,
                    },
                )
                for token in paragraph_data.get("placeholders") or []:
                    if token not in all_placeholders:
                        all_placeholders.append(token)
                if style_name.lower().startswith("heading"):
                    try:
                        level = int(style_name.split()[-1])
                    except (TypeError, ValueError):
                        level = 1
                    blocks.append(
                        block(
                            "heading",
                            text,
                            location=paragraph_data["location"],
                            level=level,
                            style=style_name,
                            placeholders=paragraph_data.get("placeholders", []),
                        )
                    )
                else:
                    blocks.append(
                        block(
                            "paragraph",
                            text,
                            location=paragraph_data["location"],
                            style=style_name,
                            placeholders=paragraph_data.get("placeholders", []),
                        )
                    )
            elif child.tag.endswith("}tbl") and options.extract_tables:
                table = Table(child, document)
                table_data = _table_payload(table, table_index + 1)
                if table_data:
                    table_index += 1
                    all_placeholders.extend(
                        token
                        for token in table_data["placeholders"]
                        if token not in all_placeholders
                    )
                    rows = table_data["rows"]
                    blocks.append(
                        block(
                            "table",
                            "\n".join(" | ".join(row) for row in rows),
                            location={
                                "part": "body",
                                "table": table_index,
                                "row_count": table_data["row_count"],
                                "column_count": table_data["column_count"],
                            },
                            columns=rows[0],
                            rows=rows[1:],
                            cells=table_data["cells"],
                            row_count=table_data["row_count"],
                            column_count=table_data["column_count"],
                            placeholders=table_data["placeholders"],
                        )
                    )

        for section_index, section in enumerate(document.sections, start=1):
            for part, paragraphs, tables in (
                ("header", section.header.paragraphs, section.header.tables),
                ("footer", section.footer.paragraphs, section.footer.tables),
            ):
                if len(blocks) >= options.max_blocks:
                    warnings.append(f"Block limit reached: {options.max_blocks}")
                    break
                for paragraph_index_in_part, paragraph in enumerate(
                    paragraphs, start=1
                ):
                    if len(blocks) >= options.max_blocks:
                        warnings.append(f"Block limit reached: {options.max_blocks}")
                        break
                    text = paragraph.text.strip()
                    if not text:
                        continue
                    paragraph_data = _paragraph_payload(
                        paragraph,
                        location={
                            "part": part,
                            "section": section_index,
                            "paragraph": paragraph_index_in_part,
                        },
                    )
                    for token in paragraph_data.get("placeholders") or []:
                        if token not in all_placeholders:
                            all_placeholders.append(token)
                    blocks.append(
                        block(
                            "header_footer",
                            text,
                            location=paragraph_data["location"],
                            style=paragraph_data["style"],
                            placeholders=paragraph_data.get("placeholders", []),
                        )
                    )
                if not options.extract_tables:
                    continue
                for table_index_in_part, table in enumerate(tables, start=1):
                    if len(blocks) >= options.max_blocks:
                        warnings.append(f"Block limit reached: {options.max_blocks}")
                        break
                    table_data = _table_payload(
                        table, table_index_in_part, part=part
                    )
                    if not table_data:
                        continue
                    for token in table_data["placeholders"]:
                        if token not in all_placeholders:
                            all_placeholders.append(token)
                    rows = table_data["rows"]
                    blocks.append(
                        block(
                            "table",
                            "\n".join(" | ".join(row) for row in rows),
                            location={
                                "part": part,
                                "section": section_index,
                                "table": table_index_in_part,
                                "row_count": table_data["row_count"],
                                "column_count": table_data["column_count"],
                            },
                            columns=rows[0],
                            rows=rows[1:],
                            cells=table_data["cells"],
                            row_count=table_data["row_count"],
                            column_count=table_data["column_count"],
                            placeholders=table_data["placeholders"],
                        )
                    )

        metadata: Dict[str, Any] = {
            "section_count": len(document.sections),
            "body_paragraph_count": paragraph_index,
            "body_table_count": table_index,
            "placeholder_count": len(all_placeholders),
            "placeholders": all_placeholders,
        }
        if document.core_properties:
            metadata.update(
                {
                    "title": document.core_properties.title,
                    "subject": document.core_properties.subject,
                    "author": document.core_properties.author,
                }
            )
        return finalize_result(
            attachment_id=attachment_id,
            path=path,
            format_name="docx",
            mime_type=mime_type,
            parser=self.name,
            blocks=blocks,
            metadata=metadata,
            warnings=warnings,
            options=options,
        )
